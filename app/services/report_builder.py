"""T3Q 본문 생성 결과(목차 트리 + 섹션별 마크다운) → hwpx/docx 조립.

설계: docs/t3q_upgrade_design.md §3. 섹션 트리는 API-RPT-001/002의
`{"name", "content"?, "references"?, "children": [...]}` 재귀 구조를 따른다.

마크다운 최소 변환 (실측: RPT-002 content는 문단 + `|` 그리드 표):
- 빈 줄로 구분된 블록 → 문단 (블록 내 단일 개행은 공백으로 합침)
- 연속된 `|` 시작 줄 → 표 (구분선 `|:---|` 행은 제거), 셀은 `|` 분리
- `**굵게**`·`# 헤딩` 기호는 텍스트만 남긴다 (서식 매핑은 범위 외)
- 그 외 기호(개조식 -, ○ 등)는 공문서 표기이므로 그대로 보존
"""
import re
from pathlib import Path

from docx import Document
from hwpx.document import HwpxDocument

# 표 구분선 행: | :--- | ---: | 형태
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:\-|]+\|?\s*$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+")


def _clean_text(text: str) -> str:
    """마크다운 강조·헤딩 기호를 제거하고 텍스트만 남긴다."""
    text = _BOLD_RE.sub(r"\1", text)
    return _HEADING_RE.sub("", text).strip()


def _parse_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [_clean_text(c) for c in cells]


def markdown_blocks(content: str) -> list[tuple[str, object]]:
    """마크다운 content를 [("p", 문단텍스트) | ("table", rows)] 블록으로 나눈다."""
    blocks: list[tuple[str, object]] = []
    para_lines: list[str] = []
    table_rows: list[list[str]] = []

    def _flush_para() -> None:
        nonlocal para_lines
        if para_lines:
            text = _clean_text(" ".join(para_lines))
            if text:
                blocks.append(("p", text))
            para_lines = []

    def _flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            blocks.append(("table", table_rows))
            table_rows = []

    for raw in (content or "").splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("|"):
            _flush_para()
            if not _TABLE_SEP_RE.match(stripped):
                table_rows.append(_parse_table_row(stripped))
            continue
        _flush_table()
        if not stripped:
            _flush_para()
        else:
            para_lines.append(stripped)
    _flush_para()
    _flush_table()
    return blocks


def flatten_leaf_names(sections: list[dict]) -> list[str]:
    """목차 트리의 리프 섹션 이름을 문서 순서로 나열한다 (진행률 계산용).

    RPT-002 스트림은 리프 섹션 단위로 결과를 보낸다 (명세·실측 일치).
    """
    leaves: list[str] = []

    def _walk(nodes: list[dict]) -> None:
        for node in nodes or []:
            children = node.get("children") or []
            if children:
                _walk(children)
            else:
                leaves.append(str(node.get("name") or ""))

    _walk(sections)
    return leaves


def _references_line(references: list[dict]) -> str | None:
    """참조문서 목록을 한 줄 표기로 만든다 (없으면 None)."""
    parts = []
    seen = set()
    for ref in references or []:
        name = str(ref.get("fileName") or "").strip()
        page = str(ref.get("page") or "").strip()
        if not name:
            continue
        key = (name, page)
        if key in seen:
            continue
        seen.add(key)
        parts.append(f"{name}(p.{page})" if page else name)
    return "※ 참고: " + ", ".join(parts) if parts else None


def _walk_sections(sections: list[dict], emit_heading, emit_para, emit_table, depth=0):
    """섹션 트리를 순회하며 헤딩·본문 블록·참조를 콜백으로 내보낸다."""
    for node in sections or []:
        name = str(node.get("name") or "").strip()
        if name:
            emit_heading(name, depth)
        for kind, payload in markdown_blocks(str(node.get("content") or "")):
            if kind == "p":
                emit_para(payload)
            else:
                emit_table(payload)
        ref_line = _references_line(node.get("references") or [])
        if ref_line:
            emit_para(ref_line)
        _walk_sections(node.get("children") or [], emit_heading, emit_para, emit_table, depth + 1)


def build_report_hwpx(title: str, sections: list[dict], output_path: str | Path) -> None:
    """목차 트리(내용 포함)를 hwpx 문서로 조립한다.

    build_hwpx(spec)는 문단·표를 그룹으로만 받아 섹션 순서를 못 지키므로,
    HwpxDocument에 직접 순서대로 쓴다 (서식 수준은 텍스트 문단 — 1차 범위).
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = HwpxDocument.new()
    try:
        if title:
            doc.add_paragraph(str(title))

        def _table(rows: list[list[str]]) -> None:
            if not rows:
                return
            n_cols = max(len(r) for r in rows)
            if n_cols == 0:
                return
            table = doc.add_table(len(rows), n_cols)
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).set_text(str(value))

        _walk_sections(
            sections,
            emit_heading=lambda name, _d: doc.add_paragraph(name),
            emit_para=lambda text: doc.add_paragraph(text),
            emit_table=_table,
        )
        doc.save_to_path(str(output_path))
    finally:
        doc.close()


def build_report_docx(title: str, sections: list[dict], output_path: str | Path) -> None:
    """목차 트리(내용 포함)를 docx 문서로 조립한다 (헤딩 레벨 반영)."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    if title:
        doc.add_heading(str(title), level=0)

    def _table(rows: list[list[str]]) -> None:
        if not rows:
            return
        n_cols = max(len(r) for r in rows)
        if n_cols == 0:
            return
        table = doc.add_table(rows=len(rows), cols=n_cols)
        table.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = str(value)

    _walk_sections(
        sections,
        emit_heading=lambda name, d: doc.add_heading(name, level=min(d + 1, 9)),
        emit_para=lambda text: doc.add_paragraph(text),
        emit_table=_table,
    )
    doc.save(str(output_path))
