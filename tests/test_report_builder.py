"""report_builder(마크다운→hwpx/docx 조립) 단위 테스트."""
from docx import Document

from app.core.hwpx import extract_hwpx, find_section_files, parse_section, validate_hwpx
from app.services.report_builder import (
    build_report_docx,
    build_report_hwpx,
    flatten_leaf_names,
    markdown_blocks,
)

SECTIONS = [
    {"name": "1. 개요", "content": "", "references": [], "children": [
        {"name": "1.1. 목적", "children": [],
         "content": "본 계획은 **신속한** 대응 체계를 구축하기 위해 수립되었다.\n\n둘째 문단이다.",
         "references": [{"id": "c1", "fileId": "f1", "fileName": "감염병예방법.pdf", "page": "4"}]},
        {"name": "1.2. 현황", "children": [],
         "content": (
             "발생 현황은 다음과 같다.\n"
             "| 구분 | 건수 |\n"
             "| :--- | ---: |\n"
             "| 1분기 | 12,450 |\n"
             "| 2분기 | 8,720 |"
         ),
         "references": []},
    ]},
    {"name": "2. 대응 체계", "children": [], "content": "부처별 대응 체계를 정비한다.", "references": []},
]


class TestMarkdownBlocks:
    def test_paragraphs_split_by_blank_line(self):
        blocks = markdown_blocks("첫 문단 첫 줄\n둘째 줄\n\n둘째 문단")
        assert blocks == [("p", "첫 문단 첫 줄 둘째 줄"), ("p", "둘째 문단")]

    def test_table_with_separator_row(self):
        blocks = markdown_blocks("| 구분 | 값 |\n| :--- | :---: |\n| A | 1 |")
        assert blocks == [("table", [["구분", "값"], ["A", "1"]])]

    def test_mixed_paragraph_table_paragraph(self):
        blocks = markdown_blocks("앞 문단\n| a | b |\n| 1 | 2 |\n뒷 문단")
        assert [k for k, _ in blocks] == ["p", "table", "p"]

    def test_bold_and_heading_markers_stripped(self):
        blocks = markdown_blocks("## 제목처럼 보이는 줄\n**강조** 텍스트")
        assert blocks == [("p", "제목처럼 보이는 줄 강조 텍스트")]

    def test_bullet_symbols_preserved(self):
        """개조식 기호(○, -)는 공문서 표기이므로 남긴다."""
        blocks = markdown_blocks("○ 첫 항목\n\n- 세부 항목")
        assert blocks == [("p", "○ 첫 항목"), ("p", "- 세부 항목")]

    def test_empty_content(self):
        assert markdown_blocks("") == []
        assert markdown_blocks(None) == []


class TestFlattenLeafNames:
    def test_nested_tree_leaves_in_order(self):
        assert flatten_leaf_names(SECTIONS) == ["1.1. 목적", "1.2. 현황", "2. 대응 체계"]

    def test_empty(self):
        assert flatten_leaf_names([]) == []


class TestBuildHwpx:
    def test_roundtrip_headings_content_and_table(self, tmp_path):
        out = tmp_path / "report.hwpx"
        build_report_hwpx("코로나19 재유행 대비계획서", SECTIONS, out)
        assert validate_hwpx(out).ok

        extract_dir = tmp_path / "x"
        extract_hwpx(out, extract_dir)
        nodes, *_ = parse_section(find_section_files(extract_dir)[0])
        texts = [n.text for n in nodes]

        assert "코로나19 재유행 대비계획서" in texts     # 제목
        assert "1.1. 목적" in texts                      # 헤딩
        assert "둘째 문단이다." in texts                  # 문단 분리
        assert any("신속한 대응 체계" in t for t in texts)  # ** 제거
        assert "12,450" in texts                          # 표 셀
        assert any(t.startswith("※ 참고:") and "감염병예방법.pdf" in t for t in texts)


class TestBuildDocx:
    def test_headings_levels_and_table(self, tmp_path):
        out = tmp_path / "report.docx"
        build_report_docx("제목", SECTIONS, out)
        doc = Document(str(out))
        para_texts = [p.text for p in doc.paragraphs]
        assert "1. 개요" in para_texts
        assert "1.1. 목적" in para_texts
        assert doc.tables and doc.tables[0].cell(1, 0).text == "1분기"
        # 헤딩 레벨: 상위 1, 하위 2
        styles = {p.text: p.style.name for p in doc.paragraphs}
        assert styles["1. 개요"] == "Heading 1"
        assert styles["1.1. 목적"] == "Heading 2"
